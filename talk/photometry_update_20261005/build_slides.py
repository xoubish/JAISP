"""Build a simple six-slide photometry update from the original PNGs."""
from pathlib import Path

import fitz
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
SLIDES = [
    (
        "output.png",
        "A simple CNN works",
        "Aperture + CNN keeps VIS residuals close to MER across most of the tested range.",
        "2,601 common held-out sources. Shading is the central 68% of residuals, not uncertainty on the median. MER is a catalog reference, not known truth.",
    ),
    (
        "PSF.png",
        "Use local PSFs",
        "Euclid archive PSFs retain structure and extended wings beyond a Gaussian approximation.",
        "These are example local Euclid Q1 GRID-PSFs. The enclosed-flux curves are normalized within each finite stamp.",
    ),
    (
        "MER1.png",
        "Agreement with MER on real sources",
        "Median offsets are small over much of the range; scatter and systematic shifts grow at the faint end.",
        "Detection-head sources on real Euclid Q1 tiles. Magnitude comparisons use positive-flux pairs. MER agreement alone does not establish flux accuracy.",
    ),
    (
        "MER_dist.png",
        "Crowding and galaxy size matter",
        "Close neighbors produce the largest offsets; extended galaxies also show size-dependent trends.",
        "This diagnostic uses MER < 24.5 sources. The stellar samples are small (10–16 objects per band). Shading spans the 16th–84th percentiles.",
    ),
    (
        "tractor1.png",
        "Known-flux injections test accuracy",
        "Foundation improves faint Euclid recovery over this Tractor baseline. Oracle fits the exact injected shapes.",
        "Real-galaxy reconstructions injected into real sky: 1,000 scenes, 100 candidate donors; 920 qualified central sources are shown. Foundation, adapted upstream Tractor and oracle fit identical pixels, positions and PSFs. Oracle knows morphology but fits flux. Shading is the within-bin 16th–84th percentile distribution. Rubin PSFs are approximate; source shot noise is absent in this pilot.",
    ),
    (
        "photrmse.png",
        "Blending remains the hard case",
        "Foundation has lower error in most faint bins; bright neighbors raise errors for both methods.",
        "VIS flux RMSE relative to known-template noise error; lower is better. This is the empirical injection pilot, not a general claim of beating Tractor or MER. The x-axis is target isolated VIS S/N. Some high-S/N blend bins favor Tractor.",
    ),
]


def text_box(slide, text, x, y, w, h, size, bold=False, color=(35, 35, 35)):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = 0
    tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = text
    p.font.name = "Arial"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = RGBColor(*color)
    return shape


def main():
    width, height = 13.333333, 7.5
    presentation = Presentation()
    presentation.slide_width = Inches(width)
    presentation.slide_height = Inches(height)
    presentation.core_properties.title = "JAISP photometry update"
    presentation.core_properties.subject = "CNN, PSFs, MER comparisons and known-flux injections"
    pdf = fitz.open()

    for number, (filename, title, description, notes) in enumerate(SLIDES, 1):
        image_path = ROOT / filename
        with Image.open(image_path) as im:
            image_width, image_height = im.size
        # Keep each complete original plot, including its labels and legend.
        box_x, box_y, box_w, box_h = 0.4, 1.05, width - 0.8, 5.55
        scale = min(box_w / image_width, box_h / image_height)
        iw, ih = image_width * scale, image_height * scale
        ix = box_x + (box_w - iw) / 2
        iy = box_y + (box_h - ih) / 2

        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor(255, 255, 255)
        text_box(slide, title, 0.45, 0.25, width - 0.9, 0.65, 30, bold=True)
        slide.shapes.add_picture(str(image_path), Inches(ix), Inches(iy), Inches(iw), Inches(ih))
        text_box(slide, description, 0.45, 6.75, width - 1.15, 0.55, 18)
        text_box(slide, str(number), width - 0.55, 7.08, 0.3, 0.2, 10, color=(120, 120, 120))
        slide.notes_slide.notes_text_frame.text = notes

        page = pdf.new_page(width=width * 72, height=height * 72)
        page.insert_text((0.45 * 72, 0.25 * 72 + 30), title, fontname="hebo", fontsize=30, color=(0.137, 0.137, 0.137))
        page.insert_image(fitz.Rect(ix * 72, iy * 72, (ix + iw) * 72, (iy + ih) * 72), filename=str(image_path))
        remaining = page.insert_textbox(
            fitz.Rect(0.45 * 72, 6.75 * 72, (width - 0.7) * 72, 7.3 * 72),
            description, fontname="helv", fontsize=18, color=(0.137, 0.137, 0.137),
        )
        if remaining < 0:
            raise ValueError(f"Caption does not fit on slide {number}")
        page.insert_text(((width - 0.5) * 72, 7.3 * 72), str(number), fontsize=10, color=(0.47, 0.47, 0.47))

    pptx_path = ROOT / "JAISP_photometry_slides.pptx"
    pdf_path = ROOT / "JAISP_photometry_slides.pdf"
    presentation.save(pptx_path)
    pdf.save(pdf_path, deflate=True)
    pdf.close()
    print(pptx_path)
    print(pdf_path)


if __name__ == "__main__":
    main()
